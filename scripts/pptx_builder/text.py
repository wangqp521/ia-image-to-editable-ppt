from __future__ import annotations

from typing import Any

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from lib.error_codes import ToolError
from .common import ObjectRegistry, rgb
from .ooxml import set_native_bullet, set_run_character_properties


ALIGNMENTS = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}
VERTICAL = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}


def _style_run(run: Any, contract: dict[str, Any], font_name: str, path: str) -> None:
    run.font.name = font_name
    run.font.size = Pt(contract["font_size"])
    run.font.bold = contract.get("font_weight", 400) >= 600
    run.font.color.rgb = rgb(contract.get("color", "#000000"), f"{path}.color")
    decoration = contract.get("decoration", "none")
    if decoration not in {"none", "underline"}:
        raise ToolError("UNSUPPORTED_FEATURE", f"{path}.decoration", decoration)
    run.font.underline = decoration == "underline"
    set_run_character_properties(run, contract, path)


def add_text_element(
    slide: Any,
    element: dict[str, Any],
    typography_item: dict[str, Any],
    registry: ObjectRegistry,
) -> None:
    element_id = element["element_id"]
    box = typography_item["text_box"]
    shape = slide.shapes.add_textbox(box["x"], box["y"], box["w"], box["h"])
    frame = shape.text_frame
    frame.clear()
    margins = box.get("margins", {})
    frame.margin_left = int(margins.get("left", 0))
    frame.margin_right = int(margins.get("right", 0))
    frame.margin_top = int(margins.get("top", 0))
    frame.margin_bottom = int(margins.get("bottom", 0))
    frame.word_wrap = bool(box.get("wrap", False))
    vertical = box.get("vertical_alignment", "top")
    if vertical not in VERTICAL:
        raise ToolError("UNSUPPORTED_FEATURE", f"typography.{element_id}.vertical_alignment", vertical)
    frame.vertical_anchor = VERTICAL[vertical]

    text = typography_item["text"]
    paragraphs = typography_item["paragraphs"]
    runs = typography_item["runs"]
    font_name = typography_item["selected_font"]
    for paragraph_index, paragraph_contract in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if paragraph_index == 0 else frame.add_paragraph()
        start, end = paragraph_contract["start"], paragraph_contract["end"]
        alignment = paragraph_contract.get("alignment", "left")
        if alignment not in ALIGNMENTS:
            raise ToolError("UNSUPPORTED_FEATURE", f"typography.{element_id}.alignment", alignment)
        paragraph.alignment = ALIGNMENTS[alignment]
        paragraph.level = int(paragraph_contract.get("list", {}).get("level", 0))
        paragraph.line_spacing = paragraph_contract.get("line_spacing", 1.0)
        paragraph.space_before = Pt(paragraph_contract.get("space_before", 0))
        paragraph.space_after = Pt(paragraph_contract.get("space_after", 0))
        if "margin_left" in paragraph_contract:
            paragraph.margin_left = int(paragraph_contract["margin_left"])
        if "indent" in paragraph_contract:
            paragraph.indent = int(paragraph_contract["indent"])
        set_native_bullet(
            paragraph,
            paragraph_contract.get("list", {"is_list": False}),
            f"modules.typography.items.{element_id}.paragraphs[{paragraph_index}].list",
        )
        for run_index, run_contract in enumerate(runs):
            run_start = max(start, run_contract["start"])
            run_end = min(end, run_contract["end"])
            if run_start >= run_end:
                continue
            run = paragraph.add_run()
            run.text = text[run_start:run_end]
            _style_run(
                run,
                run_contract,
                font_name,
                f"modules.typography.items.{element_id}.runs[{run_index}]",
            )
    registry.register(
        element_id,
        shape,
        "text",
        text_summary=text,
        font_declarations=[typography_item["internal_font_declaration"]],
    )
